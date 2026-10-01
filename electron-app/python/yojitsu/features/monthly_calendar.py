"""月間カレンダー

週ごとに「日付行＋予定行」の2行で1か月を並べ、各日のセルに「表示名：行先」を列挙する。
対象月と翌月の2シートを 月間カレンダー.xlsx に出力する。

テンプレート: シート 'yyyymm'。B〜H 列が月〜日、2行目から日付行/予定行が交互に並ぶ。
"""

import calendar
import logging
from datetime import date
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from ..context import RunContext
from ..dates import month_range, next_month, prev_month, to_yyyymm
from ..excel import copy_sheet_layout, open_or_create, order_month_sheets, replace_sheet, save_with_fallback, solid
from ..holidays import HolidayCalendar
from ..members import Member, sort_by_group
from ..schedule import DayEntry
from ..settings import DEFAULT_MONTHLY_TEMPLATE, template_file_name

logger = logging.getLogger(__name__)

OUTPUT_FILE = '月間カレンダー.xlsx'
TEMPLATE_SHEET = 'yyyymm'
FIRST_COL = 2           # B=月曜
FIRST_DATE_ROW = 2
HOLIDAY_FILL = solid('FFCCCC')
PC_CARRY_FILL = solid('FFFFE0')
PC_MARK = '★'


def _entry_text(entry: DayEntry, holidays: HolidayCalendar) -> str:
    """セルに載せる行先。表示しない日は空文字"""
    if entry.is_vacation:
        # 土日祝の「休暇」は予定なしと同じ扱い
        return '' if not holidays.is_business_day(entry.day) else '休暇'
    return entry.location


def collect_day_lines(ctx: RunContext, members: List[Member], year: int, month: int,
                      holidays: HolidayCalendar) -> Dict[int, List[Tuple[Member, str, bool]]]:
    """日 -> [(メンバ, 行先, PC持出)]（並び順はメンバ順）"""
    by_day: Dict[int, List[Tuple[Member, str, bool]]] = {}
    for member in members:
        if not member.in_monthly or not member.file_name or not ctx.schedules.has_file(member.file_name):
            continue
        schedule = ctx.schedules.month(member.file_name, year, month)
        lines = []
        for entry in (schedule or {}).values():
            text = _entry_text(entry, holidays)
            if text:
                lines.append((entry.day.day, text, bool(entry.pc)))
        if not lines:
            if member.is_proprietary:
                ctx.warn(f"{year}年{month}月の予定がありません", member.display_name)
            continue
        for day, text, has_pc in lines:
            by_day.setdefault(day, []).append((member, text, has_pc))
    return by_day


def _fill_dates(ws, year: int, month: int, weeks: List[List[int]], holidays: HolidayCalendar) -> None:
    """当月の日付と、先頭週・最終週にかかる前月末/翌月初の日付を埋める"""
    py, pm = prev_month(year, month)
    ny, nm = next_month(year, month)
    prev_last = calendar.monthrange(py, pm)[1]

    for week_index, week in enumerate(weeks):
        row = FIRST_DATE_ROW + week_index * 2
        leading_blanks = week.index(next(d for d in week if d)) if any(week) else 7
        for weekday, day in enumerate(week):
            if day:
                target = date(year, month, day)
            elif week_index == 0:
                target = date(py, pm, prev_last - (leading_blanks - weekday - 1))
            else:
                target = date(ny, nm, weekday - max(i for i, d in enumerate(week) if d))
            cell = ws.cell(row, FIRST_COL + weekday, target.day)
            # 日曜と祝日は薄赤（土曜の祝日も赤）
            if weekday == 6 or holidays.is_day_off(target):
                cell.fill = HOLIDAY_FILL


def build_sheet(ctx: RunContext, wb, template_ws, year: int, month: int) -> None:
    start, end = month_range(year, month)
    # 前月末・翌月初の色付けにも使うため前後1か月分を読む
    holidays = ctx.holidays(date(*prev_month(year, month), 1), month_range(*next_month(year, month))[1])
    members = sort_by_group(ctx.members(start, end), ctx.site)

    ws = replace_sheet(wb, to_yyyymm(year, month))
    copy_sheet_layout(template_ws, ws)
    weeks = calendar.monthcalendar(year, month)
    _fill_dates(ws, year, month, weeks, holidays)

    by_day = collect_day_lines(ctx, members, year, month, holidays)
    for week_index, week in enumerate(weeks):
        row = FIRST_DATE_ROW + week_index * 2 + 1
        for weekday, day in enumerate(week):
            lines = by_day.get(day) if day else None
            if not lines:
                continue
            # BP の PC持出だけを強調する（プロパーは対象外）
            text = "\n".join(
                f"{m.display_name}{PC_MARK if pc and not m.is_proprietary else ''}：{loc}" for m, loc, pc in lines
            )
            cell = ws.cell(row, FIRST_COL + weekday, text)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if any(pc and not m.is_proprietary for m, _, pc in lines):
                cell.fill = PC_CARRY_FILL


def run(ctx: RunContext, year: int, month: int) -> dict:
    template_path = ctx.paths.template(template_file_name(ctx.db, 'monthly_calendar', DEFAULT_MONTHLY_TEMPLATE, ctx.project_id))
    if not template_path.exists():
        raise FileNotFoundError(f"月間カレンダーのテンプレートが見つかりません: {template_path}")

    template_wb = load_workbook(template_path)
    try:
        template_ws = template_wb[TEMPLATE_SHEET]
        output_path: Path = ctx.output_dir() / OUTPUT_FILE
        wb = open_or_create(output_path)
        for y, m in ((year, month), next_month(year, month)):
            logger.info(f"月間カレンダー作成: {y}年{m}月")
            build_sheet(ctx, wb, template_ws, y, m)
        order_month_sheets(wb, to_yyyymm(year, month))
        saved = save_with_fallback(wb, output_path)
        if saved != output_path:
            ctx.warn(f"出力ファイルが開かれていたため別名で保存しました: {saved.name}")
    finally:
        template_wb.close()
    return {'outputs': [str(saved)]}
