"""個別予定ファイル（メンバごとの .xlsx）の読み込み

旧実装は機能ごとに別々の読み込み処理を持ち、列の解釈も食い違っていた。
ここで1回だけ読み、全機能が同じ DayEntry を使う。

シート構成（シート名 = YYYYMM）:
  1行目: タイトル / 2行目: ヘッダ / 3行目以降: 1日1行
  A 日付 | B 勤怠 | C 行先 | D PC持出 | E 宿泊 | F wifi |
  G〜K 実績（外部設計/内部設計/製造・単体テスト/会議/その他） | L PJ外作業 | M 備考
"""

import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, Optional, Tuple

from openpyxl import load_workbook

from .dates import to_yyyymm

logger = logging.getLogger(__name__)

FIRST_DATA_ROW = 3
COL_DATE, COL_ATTENDANCE, COL_LOCATION, COL_PC, COL_STAY, COL_WIFI = 1, 2, 3, 4, 5, 6
COL_HOURS = (7, 8, 9, 10, 11)
COL_PJ_OUTSIDE = 12
COL_REMARKS = 13

HOUR_CATEGORIES = ('外部設計', '内部設計', '製造/単体テスト', '会議', 'その他')
VACATION = '休暇'

_DATE_FORMULA = re.compile(r"=A(\d+)\+(\d+)", re.IGNORECASE)


@dataclass(frozen=True)
class DayEntry:
    day: date
    attendance: str = ''
    location: str = ''
    pc: str = ''
    stay: str = ''
    wifi: str = ''
    hours: Tuple[float, ...] = (0.0,) * len(COL_HOURS)
    pj_outside: float = 0.0
    remarks: str = ''

    @property
    def is_vacation(self) -> bool:
        return VACATION in self.attendance

    @property
    def project_hours(self) -> float:
        return sum(self.hours)

    @property
    def has_actual(self) -> bool:
        return any(h > 0 for h in self.hours) or self.pj_outside > 0


MonthSchedule = Dict[date, DayEntry]


def _text(value) -> str:
    return str(value).strip() if value is not None else ''


def _number(value) -> float:
    if value is None or value == '':
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _resolve_date(value, formula, row: int, year: int, month: int,
                  previous: Optional[Tuple[int, date]]) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 1 <= int(value) <= 31:
        return date(year, month, int(value))
    if isinstance(value, str) and value.strip().isdigit():
        return date(year, month, int(value.strip()))
    # openpyxl で作られ Excel で一度も開かれていないファイルは数式の計算結果が無い。
    # 日付列によくある =A3+1 形式だけ自前で解釈する
    if isinstance(formula, str) and previous:
        match = _DATE_FORMULA.fullmatch(formula.strip().replace('$', ''))
        prev_row, prev_date = previous
        if match and int(match.group(1)) == prev_row:
            return prev_date + timedelta(days=int(match.group(2)))
    return None


def read_month(path: Path, year: int, month: int) -> Optional[MonthSchedule]:
    """対象月シートを読む。シートが無ければ None"""
    sheet_name = to_yyyymm(year, month)
    values_wb = load_workbook(path, data_only=True, read_only=True)
    formulas_wb = load_workbook(path, data_only=False, read_only=True)
    try:
        if sheet_name not in values_wb.sheetnames:
            return None
        value_rows = values_wb[sheet_name].iter_rows(min_row=FIRST_DATA_ROW, max_col=COL_REMARKS, values_only=True)
        formula_rows = formulas_wb[sheet_name].iter_rows(min_row=FIRST_DATA_ROW, max_col=1, values_only=True)

        entries: MonthSchedule = {}
        previous: Optional[Tuple[int, date]] = None
        for row_no, (cells, formula_cells) in enumerate(zip(value_rows, formula_rows), start=FIRST_DATA_ROW):
            cells = tuple(cells) + (None,) * (COL_REMARKS - len(cells))
            day = _resolve_date(cells[COL_DATE - 1], formula_cells[0], row_no, year, month, previous)
            if day is None:
                continue
            previous = (row_no, day)
            if (day.year, day.month) != (year, month):
                continue
            entries[day] = DayEntry(
                day=day,
                attendance=_text(cells[COL_ATTENDANCE - 1]),
                location=_text(cells[COL_LOCATION - 1]),
                pc=_text(cells[COL_PC - 1]),
                stay=_text(cells[COL_STAY - 1]),
                wifi=_text(cells[COL_WIFI - 1]),
                hours=tuple(_number(cells[col - 1]) for col in COL_HOURS),
                pj_outside=_number(cells[COL_PJ_OUTSIDE - 1]),
                remarks=_text(cells[COL_REMARKS - 1]),
            )
        return entries
    finally:
        values_wb.close()
        formulas_wb.close()


class ScheduleStore:
    """個別予定ディレクトリ。ファイル名（拡張子なし）= メンバの assignment_name"""

    def __init__(self, directory: Path):
        self.directory = Path(directory)
        self._files = {p.stem: p for p in self.directory.glob('*.xlsx') if not p.name.startswith('~$')} \
            if self.directory.exists() else {}
        self._cache: Dict[Tuple[str, int, int], Optional[MonthSchedule]] = {}

    def has_file(self, file_name: str) -> bool:
        return file_name in self._files

    def month(self, file_name: str, year: int, month: int) -> Optional[MonthSchedule]:
        """ファイルもしくはシートが無ければ None。読み込み失敗は警告して None"""
        key = (file_name, year, month)
        if key not in self._cache:
            path = self._files.get(file_name)
            result = None
            if path:
                try:
                    result = read_month(path, year, month)
                except Exception as e:  # 壊れたファイル1件で全体を止めない
                    logger.warning(f"個別予定の読み込みに失敗: {path.name} ({e})")
            self._cache[key] = result
        return self._cache[key]
