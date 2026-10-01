"""XML 直接編集で、画像などテンプレートの他の要素が壊れないこと"""

import zipfile

import pytest
from lxml import etree
from openpyxl import load_workbook

from yojitsu.features.submit_files import TEMPLATE_SHEET, SheetData, XmlWriter
from yojitsu.xlsx_patch import XlsxPackage, XlsxPatchError

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}


@pytest.fixture
def template(sample_env):
    return sample_env / 'templates' / '作業実績表テンプレート_自社向け.xlsx'


def _zip(path):
    with zipfile.ZipFile(path) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _write(template, output, cells, text=None):
    XmlWriter().write(template, output, SheetData('10月実績', cells, text or {}))
    return _zip(output)


def test_image_and_other_parts_survive(template, tmp_path):
    before = _zip(template)
    after = _write(template, tmp_path / 'out.xlsx', {'B7': 2026, 'N7': '架空 太郎', 'F14': 7.5})

    media = [name for name in before if name.startswith('xl/media/')]
    assert media, 'サンプルテンプレートに画像が入っていない'
    for name in media:
        assert after[name] == before[name]
    # 元シートの図形定義は1バイトも変えない。複製先には別の図形パーツを作り、同じ画像を指す
    drawings = sorted(name for name in after if name.startswith('xl/drawings/drawing'))
    assert len(drawings) == 2
    assert after[drawings[0]] == before[drawings[0]]
    for name in drawings:
        rels = etree.fromstring(after[name.replace('drawings/', 'drawings/_rels/') + '.rels'])
        assert [r.get('Target') for r in rels] == ['/xl/media/image1.png'] or \
               [r.get('Target') for r in rels] == ['../media/image1.png']


def test_values_and_formats(template, tmp_path):
    output = tmp_path / 'out.xlsx'
    _write(template, output, {'B7': 2026, 'E7': 10, 'N7': '架空 太郎', 'F14': 7.5, 'B44': ''},
           text={'F10': '00123'})
    wb = load_workbook(output)
    assert wb.sheetnames == [TEMPLATE_SHEET, '10月実績']
    ws = wb['10月実績']
    assert (ws['B7'].value, ws['E7'].value, ws['N7'].value, ws['F14'].value) == (2026, 10, '架空 太郎', 7.5)
    assert ws['F10'].value == '00123' and ws['F10'].number_format == '@'
    assert ws['F45'].value == '=SUM(F14:F44)'               # 合計の数式は残る
    assert ws['B13'].value == '日' and ws['B13'].font.b     # 見出しの書式も残る
    assert len(ws._images) == 1
    assert wb[TEMPLATE_SHEET]['N7'].value is None            # 元シートは変えない
    assert '10月実績' in str(ws.print_area)                  # 印刷範囲も複製される


def test_workbook_recalculates_on_open(template, tmp_path):
    parts = _write(template, tmp_path / 'out.xlsx', {'F14': 8})
    workbook = etree.fromstring(parts['xl/workbook.xml'])
    assert workbook.find('m:calcPr', NS).get('fullCalcOnLoad') == '1'
    assert 'xl/calcChain.xml' not in parts


def test_rows_and_cells_are_inserted_in_order(template, tmp_path):
    output = tmp_path / 'out.xlsx'
    parts = _write(template, output, {'Z100': 'end', 'A3': 'top', 'C14': 1})
    sheet_part = next(n for n in parts if n.startswith('xl/worksheets/sheet') and n != 'xl/worksheets/sheet1.xml')
    sheet = etree.fromstring(parts[sheet_part])
    rows = [int(r.get('r')) for r in sheet.find('m:sheetData', NS)]
    assert rows == sorted(rows)
    row14 = sheet.find("m:sheetData/m:row[@r='14']", NS)
    cols = [c.get('r') for c in row14]
    assert cols == sorted(cols, key=lambda ref: (len(ref.rstrip('0123456789')), ref))


def test_existing_sheet_is_reused(template, tmp_path):
    first = tmp_path / 'first.xlsx'
    _write(template, first, {'F14': 1})
    package = XlsxPackage(first)
    package.set_cells('10月実績', {'F14': 2})
    package.save(tmp_path / 'second.xlsx')
    wb = load_workbook(tmp_path / 'second.xlsx')
    assert wb.sheetnames == [TEMPLATE_SHEET, '10月実績']
    assert wb['10月実績']['F14'].value == 2


def test_duplicate_sheet_name_is_rejected(template):
    package = XlsxPackage(template)
    with pytest.raises(XlsxPatchError):
        package.copy_sheet(TEMPLATE_SHEET, TEMPLATE_SHEET)
