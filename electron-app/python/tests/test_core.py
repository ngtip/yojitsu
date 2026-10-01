from datetime import date

import pytest

from yojitsu.db import Database
from yojitsu.features.hours_summary import summarize
from yojitsu.features.leave_matrix import symbol_for
from yojitsu.holidays import HolidayCalendar
from yojitsu.members import Member
from yojitsu.schedule import DEFAULT_LAYOUT, DayEntry, detect_layout, read_month
from yojitsu.storage.dummy import DummyStorage


def _member(**kw):
    values = dict(member_id='X', full_name='架空 太郎', display_name='架空', file_name='99_架空', group='拠点A',
                  is_proprietary=False, in_monthly=True, organization='', abbreviation='', order=0)
    values.update(kw)
    return Member(**values)


def test_business_days_exclude_holidays_but_not_events():
    calendar = HolidayCalendar.from_rows([
        {'category': '祝日', 'start_date': '2026-09-21'},
        {'category': '休暇', 'start_date': '2026-09-22', 'end_date': '2026-09-23'},
        {'category': 'イベント', 'start_date': '2026-09-24'},
    ])
    # 2026年9月の平日は22日。祝日・休暇の3日を除く
    assert calendar.business_days(date(2026, 9, 1), date(2026, 9, 30)) == 19
    assert calendar.is_public_holiday(date(2026, 9, 21))
    assert not calendar.is_public_holiday(date(2026, 9, 22))
    assert calendar.is_business_day(date(2026, 9, 24))


def test_formula_dates_are_resolved(sample_env):
    """Excel で開かれていない =A3+1 形式の日付列も読める"""
    remote = sample_env / 'remote' / 'sites' / 'sample' / 'Shared Documents' / '作業実績'
    by_formula = read_month(remote / '02_乙川.xlsx', 2026, 10)
    by_value = read_month(remote / '01_甲野.xlsx', 2026, 10)
    assert sorted(by_formula) == sorted(by_value)
    assert len(by_formula) == 31
    assert read_month(remote / '01_甲野.xlsx', 2030, 1) is None


def test_summarize_forecast():
    holidays = HolidayCalendar()
    entries = {
        date(2026, 10, 1): DayEntry(date(2026, 10, 1), attendance='出社', hours=(10.0, 0, 0, 0, 0)),
        date(2026, 10, 2): DayEntry(date(2026, 10, 2), attendance='休暇'),
        date(2026, 10, 3): DayEntry(date(2026, 10, 3), attendance='休日出勤', hours=(4.0, 0, 0, 0, 0)),
        date(2026, 10, 5): DayEntry(date(2026, 10, 5), attendance='出社', hours=(8.0, 0, 0, 0, 0), pj_outside=1.0),
    }
    result = summarize(_member(), entries, holidays, 2026, 10)
    assert result.business_days == 22
    assert result.actual == 22.0
    assert result.pj_outside == 1.0
    assert result.business_input_days == 3          # 1日・2日（休暇）・5日
    assert result.holiday_work_planned == 1 and result.holiday_work_actual == 1
    assert result.avg_overtime == 0.5               # 残業2h / 3日 = 0.67 → 0.5刻み
    assert result.forecast_standard == 22.0 + 19 * 8
    assert result.forecast_with_overtime == 22.0 + 19 * 8.5


@pytest.mark.parametrize('attendance, holiday, expected', [
    ('A休', False, '△'), ('P休', False, '▽'), ('休暇', False, '〇'), ('', True, '□'), ('出社', False, ''),
])
def test_leave_symbols(attendance, holiday, expected):
    assert symbol_for(attendance, holiday) == expected


def test_holiday_crud_and_dedup(db_file):
    with Database(db_file) as db:
        first = db.add_holiday('架空記念日', '祝日', '2026-12-01')
        assert db.add_holiday('架空記念日', '祝日', '2026-12-01') == first
        db.update_holiday(first, end_date='2026-12-02')
        assert db.get_holidays(None, '2026-12-02', '2026-12-02')[0]['holiday_id'] == first
        with pytest.raises(ValueError):
            db.update_holiday(first, start_date='2026-12-05')     # 終了日より後
        db.delete_holiday(first)
        with pytest.raises(ValueError):
            db.delete_holiday(first)


def test_member_update_validation(db_file):
    with Database(db_file) as db:
        updated = db.update_member('M003', display_name='丙田改')
        assert updated['display_name'] == '丙田改' and updated['full_name'] == '丙田 三郎'
        with pytest.raises(ValueError):
            db.update_member('M003', file_storage_location='cloud')
        with pytest.raises(ValueError):
            db.update_member('NOPE', display_name='x')


def test_config_overwrites_global_value(db_file):
    with Database(db_file) as db:
        db.set_config('monthly_calendar', 'template_file', 'a.xlsx')
        db.set_config('monthly_calendar', 'template_file', 'b.xlsx')
        db.set_config('monthly_calendar', 'template_file', 'pj.xlsx', project_id='PJ-SAMPLE')
        assert db.get_config('monthly_calendar', 'template_file') == 'b.xlsx'
        assert db.get_config('monthly_calendar', 'template_file', 'PJ-SAMPLE') == 'pj.xlsx'


def test_dummy_storage_rejects_parent_paths(tmp_path):
    storage = DummyStorage(tmp_path)
    with pytest.raises(ValueError):
        storage.list_files('/sites/../../etc')


def test_layout_follows_header_row():
    current = ('日付', '勤怠', '行先', 'PC持出', '入館証持出', 'wifi持出',
               '外部設計', '内部設計', '製造\n単体テスト', '会議', 'その他', 'PJ外作業', '備考')
    assert detect_layout(current) == DEFAULT_LAYOUT
    # 旧レイアウト: PJ外作業が無く L が備考
    old = detect_layout(current[:11] + ('備考',))
    assert old['remarks'] == 12 and 'pj_outside' not in old
    # 見出しが読めないシートは既定の並び
    assert detect_layout((None, None)) == DEFAULT_LAYOUT


def test_old_layout_sheet_reads_remarks_not_as_hours(sample_env):
    """旧レイアウトの月は L 列（備考）を PJ外作業として数えない"""
    remote = sample_env / 'remote' / 'sites' / 'sample' / 'Shared Documents' / '作業実績'
    oldest = read_month(remote / '01_甲野.xlsx', 2026, 9)      # 01 は最古の月が旧レイアウト
    assert oldest and all(e.pj_outside == 0 for e in oldest.values())


@pytest.mark.parametrize('value, expected', [('有', True), ('無', False), ('－', False), ('', False), ('〇', True)])
def test_pc_carry_only_when_yes(value, expected):
    assert DayEntry(date(2026, 10, 1), pc=value).carries_pc is expected


def test_group_counts_use_keywords(tmp_path):
    """拠点の表示名と行先が違う（行先に最寄り駅名を書く拠点など）場合はキーワードで数える"""
    from yojitsu.features.list_calendar import _location_counts
    from yojitsu.settings import load_site_settings

    path = tmp_path / 'site.json'
    path.write_text('{"groups": [{"name": "拠点X", "keywords": ["本社", "支所"]}, {"name": "拠点Y"}]}', encoding='utf-8')
    site = load_site_settings(path)
    day = date(2026, 10, 1)
    members = [_member(member_id=str(i)) for i in range(4)]
    schedules = {str(i): {day: DayEntry(day, location=loc)} for i, loc in enumerate(['本社', '支所', '拠点Y', '在宅'])}
    assert _location_counts(day, members, schedules, site.groups) == {'拠点X': 2, '拠点Y': 1}


def test_members_with_several_stints_are_merged(db_file):
    from yojitsu.members import load_members
    from yojitsu.settings import SiteSettings

    with Database(db_file) as db:
        members = load_members(db, 'PJ-SAMPLE', SiteSettings(), date(2026, 10, 1), date(2026, 10, 31))
    tei = [m for m in members if m.display_name == '丁村']
    assert len(tei) == 1 and len(tei[0].periods) == 2
    assert tei[0].participates(date(2026, 10, 5)) and not tei[0].participates(date(2026, 10, 15))
    assert tei[0].file_name == '04_丁村'


def test_schedule_file_name_is_migrated_from_assignment(tmp_path):
    """旧DB（members に schedule_file_name が無い）を開くと、PJ所属のファイル名を移す"""
    import sqlite3
    path = tmp_path / 'old.sqlite'
    con = sqlite3.connect(path)
    con.executescript("""
        CREATE TABLE projects (project_id TEXT PRIMARY KEY, project_name TEXT, project_code TEXT);
        CREATE TABLE members (member_id TEXT PRIMARY KEY, full_name TEXT NOT NULL);
        CREATE TABLE project_members (project_member_id INTEGER PRIMARY KEY, project_id TEXT, member_id TEXT,
                                      start_date TEXT, end_date TEXT, assignment_name TEXT);
        INSERT INTO members VALUES ('M1', '架空 太郎');
        INSERT INTO project_members VALUES (1, 'P', 'M1', '2026-01-01', '2026-03-31', '01_旧名');
        INSERT INTO project_members VALUES (2, 'P', 'M1', '2026-06-01', NULL, '01_架空');
    """)
    con.commit()
    con.close()
    with Database(path) as db:
        assert db.get_member('M1')['schedule_file_name'] == '01_架空'   # 最新の参画のもの


def test_unknown_groups_are_warned_and_kept_together(db_file, tmp_path, monkeypatch):
    """site-settings が無い／拠点名が一致しないときも拠点ごとにまとめ、画面に警告を出す"""
    from yojitsu.context import RunContext
    from yojitsu.members import sort_by_group

    monkeypatch.setenv('YOJITSU_SITE_SETTINGS', str(tmp_path / 'missing.json'))
    with Database(db_file) as db:
        ctx = RunContext(db)
        members = sort_by_group(ctx.members(date(2026, 10, 1), date(2026, 10, 31)), ctx.site)
    assert 'site-settings.json が無い' in ctx.warnings[0].message
    groups = [m.group for m in members]
    assert groups == sorted(groups, key=groups.index)          # 拠点が交互に混ざらない
    assert groups.index('拠点B') > max(i for i, g in enumerate(groups) if g == '拠点A')

    (tmp_path / 'site.json').write_text('{"groups": [{"name": "拠点B"}]}', encoding='utf-8')
    monkeypatch.setenv('YOJITSU_SITE_SETTINGS', str(tmp_path / 'site.json'))
    with Database(db_file) as db:
        ctx = RunContext(db)
        members = sort_by_group(ctx.members(date(2026, 10, 1), date(2026, 10, 31)), ctx.site)
    assert [w.message for w in ctx.warnings] == ['拠点「拠点A」が site-settings.json の groups にありません（並び順・人数集計の対象外）']
    assert members[0].group == '拠点B'                          # 設定にある拠点が先


@pytest.mark.parametrize('count, marked', [(8, False), (9, True)])
def test_group_is_marked_only_above_seat_count(tmp_path, count, marked):
    """座席数 8 の拠点は 9 人以上で黄色（8 人ちょうどは黄色にしない）"""
    from types import SimpleNamespace

    from openpyxl import Workbook

    from yojitsu.features.list_calendar import HIGHLIGHT, _Row, _write_day
    from yojitsu.settings import load_site_settings

    path = tmp_path / 'site.json'
    path.write_text('{"groups": [{"name": "拠点X", "threshold": 8}, {"name": "拠点Y", "threshold": 4}]}',
                    encoding='utf-8')
    ctx = SimpleNamespace(site=load_site_settings(path))
    day = date(2026, 10, 1)
    rows = [_Row(_member(member_id=str(i), group='拠点X'), None) for i in range(count)]
    ws = Workbook().active
    _write_day(ws, 3, day, rows, {'拠点X': count, '拠点Y': 0}, ctx, HolidayCalendar())
    assert (ws['B3'].fill.start_color.rgb == '00' + HIGHLIGHT) is marked
    assert ws['B4'].value == f'{count}名'
