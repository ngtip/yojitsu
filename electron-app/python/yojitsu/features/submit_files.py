"""作業実績表（提出用）の作成

BP: 自社向け と PJ向け の2ファイル / プロパー: PJ向け のみ
出力先: <出力Dir>/<YYYY>/<MM>/{自社向け,PJ向け}/

テンプレートのシート 'xx月実績' を複製して '<M>月実績' とし、次のセルに書き込む。
  B7 年 / E7 月 / H7 所属（自社向け=BP会社名, PJ向け=自社正式名） / N7 氏名
  F9 PJ名 / F10 PJコード（自社向けのみ・文字列）
  14〜44行目 = 1〜31日: B 日 / D 曜日 / F〜J 実績（PJ向けは F に合計のみ）

テンプレートに画像があると openpyxl では保存時に失われるため、実環境では Excel COM で書く。
COM が使えない環境（このリポジトリのテスト等）では openpyxl で同じセルに書く。
"""

import logging
import os
import shutil
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Protocol

from ..context import RunContext
from ..dates import WEEKDAY_JA, month_range, to_yyyymm
from ..members import Member
from ..schedule import DayEntry

logger = logging.getLogger(__name__)

TEMPLATE_SHEET = 'xx月実績'
FIRST_DAY_ROW = 14
LAST_DAY_ROW = 44
COL_DAY, COL_WEEKDAY = 2, 4
COL_HOURS_FIRST = 6     # F〜J
OWN, PJ = '自社向け', 'PJ向け'


@dataclass
class SheetData:
    """1ファイル分の書き込み内容（セル番地 -> 値）"""
    sheet_name: str
    cells: Dict[str, object]
    text_cells: Dict[str, str]  # 文字列書式で書くセル


def build_sheet_data(kind: str, member: Member, entries: Dict[date, DayEntry], year: int, month: int,
                     ctx: RunContext) -> SheetData:
    cells: Dict[str, object] = {'B7': year, 'E7': month, 'F9': ctx.project_name}
    text_cells: Dict[str, str] = {}
    if member.full_name:
        cells['N7'] = member.full_name
    if kind == OWN:
        if member.organization:
            cells['H7'] = member.organization
        if ctx.project_code:
            text_cells['F10'] = str(ctx.project_code)
    else:
        cells['H7'] = ctx.site.own_company_fullname

    _, last = month_range(year, month)
    for day in range(1, 32):
        row = FIRST_DAY_ROW + day - 1
        if day <= last.day:
            cells[f'B{row}'] = day
            cells[f'D{row}'] = WEEKDAY_JA[last.replace(day=day).weekday()]
        else:
            cells[f'B{row}'] = ''
            cells[f'D{row}'] = ''

    for entry in entries.values():
        row = FIRST_DAY_ROW + entry.day.day - 1
        if kind == PJ:
            if entry.project_hours > 0:
                cells[_ref(COL_HOURS_FIRST, row)] = entry.project_hours
        else:
            for offset, hours in enumerate(entry.hours):
                if hours > 0:
                    cells[_ref(COL_HOURS_FIRST + offset, row)] = hours
    return SheetData(f"{month}月実績", cells, text_cells)


def _ref(col: int, row: int) -> str:
    return f"{chr(ord('A') + col - 1)}{row}"


def file_name_for(kind: str, member: Member, ctx: RunContext, yyyymm: str) -> str:
    company = (member.abbreviation or '－') if kind == OWN else ctx.site.own_company_shortname
    return f"{ctx.project_name}（{company}_{member.display_name}）作業実績表_{yyyymm}.xlsx"


class Writer(Protocol):
    def write(self, template: Path, output: Path, data: SheetData) -> None: ...
    def close(self) -> None: ...


class OpenpyxlWriter:
    def write(self, template: Path, output: Path, data: SheetData) -> None:
        from openpyxl import load_workbook

        shutil.copy2(template, output)
        wb = load_workbook(output)
        try:
            if data.sheet_name in wb.sheetnames:
                ws = wb[data.sheet_name]
            else:
                ws = wb.copy_worksheet(wb[TEMPLATE_SHEET])
                ws.title = data.sheet_name
            for ref, value in data.cells.items():
                ws[ref] = value
            for ref, value in data.text_cells.items():
                ws[ref].number_format = '@'
                ws[ref] = value
            wb.save(output)
        finally:
            wb.close()

    def close(self) -> None:
        pass


class ComWriter:
    """Excel COM。テンプレートは種類ごとに1度だけ開き、SaveCopyAs で保存して使い回す"""

    def __init__(self):
        import win32com.client  # type: ignore

        self.excel = win32com.client.DispatchEx('Excel.Application')
        self.excel.Visible = False
        self.excel.DisplayAlerts = False
        self.workbooks = {}

    def write(self, template: Path, output: Path, data: SheetData) -> None:
        key = str(template.resolve())
        if key not in self.workbooks:
            self.workbooks[key] = self.excel.Workbooks.Open(key, ReadOnly=True)
        wb = self.workbooks[key]
        try:
            ws = wb.Worksheets(data.sheet_name)
        except Exception:
            wb.Worksheets(TEMPLATE_SHEET).Copy(After=wb.Worksheets(wb.Worksheets.Count))
            ws = wb.Worksheets(wb.Worksheets.Count)
            ws.Name = data.sheet_name
        try:
            for ref, value in data.cells.items():
                ws.Range(ref).Value = value
            for ref, value in data.text_cells.items():
                ws.Range(ref).NumberFormat = '@'
                ws.Range(ref).Value = value
            wb.SaveCopyAs(str(output.resolve()))
        finally:
            # 次のメンバに値が残らないよう書いたセルを空に戻す
            for ref in list(data.cells) + list(data.text_cells):
                ws.Range(ref).Value = ''

    def close(self) -> None:
        for wb in self.workbooks.values():
            wb.Close(SaveChanges=False)
        self.excel.Quit()


def create_writer(kind: Optional[str] = None) -> Writer:
    kind = (kind or os.environ.get('YOJITSU_SUBMIT_WRITER') or 'auto').lower()
    if kind in ('com', 'auto'):
        try:
            return ComWriter()
        except Exception as e:
            if kind == 'com':
                raise
            logger.warning(f"Excel COM が使えないため openpyxl で作成します（画像は保持されません）: {e}")
    return OpenpyxlWriter()


def run(ctx: RunContext, year: int, month: int, writer_kind: Optional[str] = None) -> dict:
    templates = {OWN: ctx.paths.template(ctx.site.own_template_file),
                 PJ: ctx.paths.template(ctx.site.pj_template_file)}
    for kind, path in templates.items():
        if not path.exists():
            raise FileNotFoundError(f"作業実績表テンプレート（{kind}）が見つかりません: {path}")

    yyyymm = to_yyyymm(year, month)
    start, end = month_range(year, month)
    outputs: List[str] = []
    writer = create_writer(writer_kind)
    try:
        for member in ctx.members(start, end):
            if not ctx.schedules.has_file(member.file_name):
                ctx.warn('個別予定ファイルが見つかりません', member.display_name)
                continue
            entries = ctx.schedules.month(member.file_name, year, month)
            if not entries:
                ctx.warn(f"{year}年{month}月のデータがありません", member.display_name)
                continue
            kinds = [PJ] if member.is_proprietary else [OWN, PJ]
            for kind in kinds:
                output = ctx.output_dir(f"{year}", f"{month:02d}", kind) / file_name_for(kind, member, ctx, yyyymm)
                writer.write(templates[kind], output, build_sheet_data(kind, member, entries, year, month, ctx))
                logger.info(f"作成: {kind}/{output.name}")
                outputs.append(str(output))
    finally:
        writer.close()
    return {'outputs': outputs}
