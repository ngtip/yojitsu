"""一覧カレンダー

1日ごとにメンバを縦に並べ、勤怠・行先・PC持出などを一覧にする。
対象月と翌月の2シートを 一覧カレンダー.xlsx に出力する。

テンプレート: シート 'yyyymm'（1〜2行目がタイトルとヘッダ。データは3行目から）
  A 日付 | B 拠点別人数 | C メンバ | D 勤怠 | E 行先 | F PC持出 | G 入館証持出 | H wifi持出 | I 備考

表示ルール:
  - 平日は全メンバ、土日祝は勤怠が入っているメンバだけ（誰もいなければ空行1行）
  - B列: 拠点の先頭行にラベル、次の行に「行先がそのラベルを含む人数」。
    人数が拠点の閾値を超えたら拠点の範囲を黄色にする
  - A列: site-settings の date_highlight に指定した拠点の人数が規定以上なら黄色
  - A列の日付は各日の1行目だけ見えるようにし、2行目以降は背景色と同じ文字色にする
"""

import logging
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.styles import Border, Font

from ..context import RunContext
from ..dates import iter_dates, month_range, next_month, to_yyyymm
from ..excel import (THIN, THIN_BORDER, copy_sheet_layout, open_or_create, order_month_sheets,
                     replace_sheet, save_with_fallback, solid)
from ..holidays import HolidayCalendar
from ..members import Member, sort_by_group
from ..schedule import DayEntry
from ..settings import DEFAULT_LIST_TEMPLATE, template_file_name

logger = logging.getLogger(__name__)

OUTPUT_FILE = '一覧カレンダー.xlsx'
TEMPLATE_SHEET = 'yyyymm'
FIRST_DATA_ROW = 3
LAST_COL = 9            # I列
NONE_MARK = '－'

WHITE, SATURDAY, HOLIDAY, HIGHLIGHT = 'FFFFFF', 'CCDDFF', 'FFCCCC', 'FFF200'
BLACK = Font(color='000000')


@dataclass
class _Row:
    member: Optional[Member]
    entry: Optional[DayEntry]


def _day_color(day: date, holidays: HolidayCalendar) -> str:
    if day.weekday() == 6 or holidays.is_day_off(day):
        return HOLIDAY          # 土曜の祝日も赤を優先
    if day.weekday() == 5:
        return SATURDAY
    return WHITE


def _rows_for_day(day: date, members: List[Member], schedules: Dict[str, Dict[date, DayEntry]],
                  holidays: HolidayCalendar) -> List[_Row]:
    business_day = holidays.is_business_day(day)
    rows = []
    for member in members:
        entry = schedules[member.member_id].get(day)
        # 土日祝の「休暇」は予定なしと同じ
        if entry and entry.is_vacation and not business_day:
            entry = None
        if business_day or (entry and entry.attendance):
            rows.append(_Row(member, entry))
    return rows or [_Row(None, None)]


def _location_counts(day: date, members: List[Member], schedules, labels: List[str]) -> Dict[str, int]:
    counts = {label: 0 for label in labels}
    for member in members:
        entry = schedules[member.member_id].get(day)
        if not entry:
            continue
        for label in labels:
            if label in entry.location:
                counts[label] += 1
    return counts


def _write_day(ws, start_row: int, day: date, rows: List[_Row], counts: Dict[str, int],
               ctx: RunContext, holidays: HolidayCalendar) -> int:
    site = ctx.site
    thresholds = {g.label: g.threshold for g in site.groups}
    day_color = _day_color(day, holidays)
    highlight_count = counts.get(site.date_highlight_label, 0)
    date_color = HIGHLIGHT if site.date_highlight_min_count and highlight_count >= site.date_highlight_min_count \
        else day_color

    # 拠点ごとの開始行（1始まりの日内行番号）と範囲
    group_starts: Dict[str, int] = {}
    for index, row in enumerate(rows, start=1):
        if row.member and row.member.group not in group_starts:
            group_starts[row.member.group] = index
    ordered = sorted(group_starts.items(), key=lambda item: item[1])
    group_spans = {
        group: (first, ordered[i + 1][1] - 1 if i + 1 < len(ordered) else len(rows))
        for i, (group, first) in enumerate(ordered)
    }

    b_values: Dict[int, str] = {}
    b_highlight = set()
    start_rows = set(group_starts.values())
    for group in site.groups:
        if group.name not in group_starts:
            continue
        first, last = group_spans[group.name]
        count_text = f"{counts.get(group.label, 0)}名"
        # 原則はラベルと人数の2行表示。次の行が別拠点の開始行なら1行に併記する
        if first + 1 <= len(rows) and first + 1 not in start_rows:
            b_values[first] = group.label
            b_values[first + 1] = count_text
        else:
            b_values[first] = f"{group.label} {count_text}"
        if counts.get(group.label, 0) > thresholds.get(group.label, 999):
            b_highlight.update(range(first, last + 1))

    total = len(rows)
    for index, row in enumerate(rows, start=1):
        excel_row = start_row + index - 1
        entry = row.entry
        values = [datetime(day.year, day.month, day.day), b_values.get(index, '')]
        if row.member:
            values += [
                row.member.display_name,
                entry.attendance if entry else '',
                entry.location if entry else '',
                (entry.pc if entry else '') or NONE_MARK,
                (entry.badge if entry else '') or NONE_MARK,
                (entry.wifi if entry else '') or NONE_MARK,
                entry.remarks if entry else '',
            ]
        else:
            values += [''] * (LAST_COL - 2)

        group = row.member.group if row.member else None
        prev_group = rows[index - 2].member.group if index > 1 and rows[index - 2].member else None
        next_group = rows[index].member.group if index < total and rows[index].member else None

        # A/B列は日付ブロックを1つの枠にする（内側の上下罫線なし）
        block_border = Border(left=THIN, right=THIN,
                              top=THIN if index == 1 else None,
                              bottom=THIN if index == total else None)

        for col, value in enumerate(values, start=1):
            cell = ws.cell(excel_row, col, value)
            if col == 1:
                cell.number_format = 'm/d(aaa)'
                cell.fill = solid(date_color)
                cell.font = BLACK if index == 1 else Font(color=date_color)
                cell.border = block_border
            elif col == 2:
                cell.fill = solid(HIGHLIGHT if index in b_highlight else day_color)
                cell.font = BLACK
                # 拠点の区切りに線を引く
                cell.border = Border(
                    left=THIN, right=THIN,
                    top=THIN if row.member and prev_group != group else block_border.top,
                    bottom=THIN if row.member and next_group != group else block_border.bottom,
                ) if row.member else block_border
            else:
                cell.fill = solid(day_color)
                cell.font = BLACK
                cell.border = THIN_BORDER
    return start_row + total


def build_sheet(ctx: RunContext, wb, template_ws, year: int, month: int) -> None:
    start, end = month_range(year, month)
    holidays = ctx.holidays(start, end)
    members = sort_by_group(ctx.members(start, end), ctx.site)

    schedules: Dict[str, Dict[date, DayEntry]] = {}
    for member in members:
        schedule = ctx.schedules.month(member.file_name, year, month) if member.file_name else None
        if schedule is None and member.is_proprietary:
            ctx.warn(f"{year}年{month}月のシートがありません", member.display_name)
        schedules[member.member_id] = schedule or {}
    # 予定ファイルが無いメンバは載せない（旧仕様どおり）
    listed = [m for m in members if ctx.schedules.has_file(m.file_name)]

    ws = replace_sheet(wb, to_yyyymm(year, month))
    copy_sheet_layout(template_ws, ws)

    labels = [g.label for g in ctx.site.groups]
    row = FIRST_DATA_ROW
    for day in iter_dates(start, end):
        counts = _location_counts(day, listed, schedules, labels)
        row = _write_day(ws, row, day, _rows_for_day(day, listed, schedules, holidays), counts, ctx, holidays)

    ws.freeze_panes = f'A{FIRST_DATA_ROW}'
    ws.auto_filter.ref = f"A{FIRST_DATA_ROW - 1}:I{row - 1}"


def run(ctx: RunContext, year: int, month: int) -> dict:
    template_path = ctx.paths.template(template_file_name(ctx.db, 'list_calendar', DEFAULT_LIST_TEMPLATE, ctx.project_id))
    if not template_path.exists():
        raise FileNotFoundError(f"一覧カレンダーのテンプレートが見つかりません: {template_path}")

    template_wb = load_workbook(template_path)
    try:
        output_path: Path = ctx.output_dir() / OUTPUT_FILE
        wb = open_or_create(output_path)
        for y, m in ((year, month), next_month(year, month)):
            logger.info(f"一覧カレンダー作成: {y}年{m}月")
            build_sheet(ctx, wb, template_wb[TEMPLATE_SHEET], y, m)
        order_month_sheets(wb, to_yyyymm(year, month))
        saved = save_with_fallback(wb, output_path)
        if saved != output_path:
            ctx.warn(f"出力ファイルが開かれていたため別名で保存しました: {saved.name}")
    finally:
        template_wb.close()
    return {'outputs': [str(saved)]}
