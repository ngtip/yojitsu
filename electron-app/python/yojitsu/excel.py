"""openpyxl の共通操作"""

import logging
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

THIN = Side(style='thin', color='000000')
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def solid(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type='solid')


def copy_sheet_layout(src: Worksheet, dst: Worksheet) -> None:
    """値・書式・列幅・行高をコピーする（別ブック間でも使える）"""
    for row in src.iter_rows():
        for cell in row:
            target = dst.cell(cell.row, cell.column)
            if cell.value is not None:
                target.value = cell.value
            if cell.has_style:
                target.font = copy(cell.font)
                target.border = copy(cell.border)
                target.fill = copy(cell.fill)
                target.number_format = cell.number_format
                target.alignment = copy(cell.alignment)
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
    for key, dim in src.row_dimensions.items():
        dst.row_dimensions[key].height = dim.height
    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))


def open_or_create(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def replace_sheet(wb: Workbook, title: str) -> Worksheet:
    """同名シートがあれば同じ位置に作り直す"""
    index = None
    if title in wb.sheetnames:
        index = wb.sheetnames.index(title)
        del wb[title]
    return wb.create_sheet(title, index)


def order_month_sheets(wb: Workbook, active_title: str) -> None:
    """YYYYMM シートを新しい月から順に並べ、指定シートをアクティブにする"""
    month_sheets = sorted((s for s in wb.sheetnames if len(s) == 6 and s.isdigit()), reverse=True)
    for index, title in enumerate(month_sheets):
        ws = wb[title]
        wb.move_sheet(ws, offset=index - wb.index(ws))
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    if active_title in wb.sheetnames:
        wb.active = wb.sheetnames.index(active_title)
        wb[active_title].sheet_view.tabSelected = True


def save_with_fallback(wb: Workbook, path: Path) -> Path:
    """出力先を Excel で開いたままでも止まらないよう、失敗時は時刻付きの別名で保存する"""
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        wb.save(alt)
        logger.warning(f"{path.name} が開かれているため別名で保存しました: {alt.name}")
        return alt
