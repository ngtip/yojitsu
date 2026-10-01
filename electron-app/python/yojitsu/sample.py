"""架空データのテスト環境を作る

  <dest>/
    db/management.sqlite        ... PJ・メンバ・祝日・設定（storage backend = dummy）
    templates/                  ... 帳票テンプレート（最小構成の見本）
    remote/sites/sample/...     ... SharePoint に見立てた個別予定の置き場
    schedules/                  ... 同期先（最初は空）
    output/                     ... 出力先
    site-settings.json          ... 自社名・拠点（架空）

人名・社名はすべて架空。実データは絶対に入れないこと。
"""

import random
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .dates import iter_dates, month_range, next_month, prev_month, to_yyyymm
from .db import Database
from .excel import THIN_BORDER, solid

PROJECT_ID = 'PJ-SAMPLE'
REMOTE_WORK_DIR = '/sites/sample/Shared Documents/作業実績'

# (member_id, 氏名, 表示名, 拠点, 所属, 略称, プロパー, 格納区分, 月間表示)
MEMBERS = [
    ('M001', '甲野 一郎', '甲野', '拠点A', '自社', '', 1, 'internal', 1),
    ('M002', '乙川 二葉', '乙川', '拠点A', '自社', '', 1, 'internal', 1),
    ('M003', '丙田 三郎', '丙田', '拠点A', '架空システム株式会社', '架空SYS', 0, 'internal', 1),
    ('M004', '丁村 四季', '丁村', '拠点B', '見本テック合同会社', '見本T', 0, 'internal', 1),
    ('M005', '戊井 五郎', '戊井', '拠点B', '試験ソフト株式会社', '試験S', 0, 'internal', 0),
    ('M006', '己島 六花', '己島', '拠点B', '外部委託株式会社', '外部委', 0, 'external', 1),
]

# 2026〜2027年の祝日（サンプル用）
PUBLIC_HOLIDAYS = [
    ('2026-01-01', '元日'), ('2026-01-12', '成人の日'), ('2026-02-11', '建国記念の日'),
    ('2026-02-23', '天皇誕生日'), ('2026-03-20', '春分の日'), ('2026-04-29', '昭和の日'),
    ('2026-05-03', '憲法記念日'), ('2026-05-04', 'みどりの日'), ('2026-05-05', 'こどもの日'),
    ('2026-05-06', '振替休日'), ('2026-07-20', '海の日'), ('2026-08-11', '山の日'),
    ('2026-09-21', '敬老の日'), ('2026-09-22', '国民の休日'), ('2026-09-23', '秋分の日'),
    ('2026-10-12', 'スポーツの日'), ('2026-11-03', '文化の日'), ('2026-11-23', '勤労感謝の日'),
    ('2027-01-01', '元日'), ('2027-01-11', '成人の日'), ('2027-02-11', '建国記念の日'),
    ('2027-02-23', '天皇誕生日'), ('2027-03-21', '春分の日'), ('2027-03-22', '振替休日'),
]
COMPANY_HOLIDAYS = [
    ('夏季休暇', '休暇', '2026-08-12', '2026-08-14'),
    ('年末年始休暇', '休暇', '2026-12-29', '2027-01-03'),
]

HEADER_FILL = solid('DDEBF7')
ACTUAL_FILL = solid('E2EFDA')
SCHEDULE_HEADERS = ('日付', '勤怠', '行先', 'PC持出', '宿泊', 'wifi',
                    '外部設計', '内部設計', '製造/単体テスト', '会議', 'その他', 'PJ外作業', '備考')


def create_sample_environment(dest: Path, base: Optional[date] = None, force: bool = False) -> List[Path]:
    dest = Path(dest).resolve()
    if dest.exists():
        if not force:
            raise FileExistsError(f"既に存在します（作り直すなら --force）: {dest}")
        shutil.rmtree(dest)
    base = base or date.today()
    months = _months_around(base)

    for sub in ('db', 'templates', 'schedules', 'output'):
        (dest / sub).mkdir(parents=True, exist_ok=True)
    remote_dir = dest / 'remote' / Path(*REMOTE_WORK_DIR.strip('/').split('/'))
    remote_dir.mkdir(parents=True)

    db_path = dest / 'db' / 'management.sqlite'
    _create_db(db_path, dest, months)
    _create_templates(dest / 'templates')
    created = [db_path]
    rng = random.Random(20261001)
    for index, member in enumerate(MEMBERS, start=1):
        member_id, _, display, group, *_rest = member
        storage = member[7]
        # 外部格納のメンバはフォルダごと分けて置く（同期で除外される）
        folder = remote_dir / f"{index:02d}_{display}" if storage == 'external' else remote_dir
        folder.mkdir(exist_ok=True)
        path = folder / f"{index:02d}_{display}.xlsx"
        _create_schedule(path, display, group, months, base, rng, use_formula_dates=(index == 2))
        created.append(path)

    (dest / 'site-settings.json').write_text(
        '{\n'
        '  "own_company_shortname": "自社",\n'
        '  "own_company_fullname": "サンプル自社株式会社",\n'
        '  "own_template_file": "作業実績表テンプレート_自社向け.xlsx",\n'
        '  "pj_template_file": "作業実績表テンプレート_PJ向け.xlsx",\n'
        '  "groups": [\n'
        '    { "name": "拠点A", "label": "拠点A", "threshold": 2 },\n'
        '    { "name": "拠点B", "label": "拠点B", "threshold": 1 }\n'
        '  ],\n'
        '  "date_highlight": { "label": "拠点A", "min_count": 3 }\n'
        '}\n',
        encoding='utf-8',
    )
    return created


def _months_around(base: date) -> List[Tuple[int, int]]:
    first = prev_month(base.year, base.month)
    months = [first]
    for _ in range(3):
        months.append(next_month(*months[-1]))
    return months


def _create_db(path: Path, root: Path, months) -> None:
    start, _ = month_range(*months[0])
    with Database(path, create=True) as db:
        db.init_tool_settings(str(root), 'templates', 'output', 'schedules')
        db.set_config('storage', 'backend', 'dummy')
        db.set_config('storage', 'dummy_root', str(root / 'remote'))
        db.set_config('sharepoint', 'remote_work_dir', REMOTE_WORK_DIR)
        db.set_config('sharepoint', 'site_url', 'https://example.sharepoint.com/sites/sample')
        with db.conn:
            db.conn.execute(
                "INSERT INTO projects (project_id, project_name, project_code, admin_name, start_date, is_active)"
                " VALUES (?, ?, ?, ?, ?, 1)",
                (PROJECT_ID, 'サンプルPJ', 'S-0001', '甲野 一郎', '2026-01-01'),
            )
            for index, (mid, full, disp, group, org, abbr, prop, storage, monthly) in enumerate(MEMBERS, start=1):
                db.conn.execute(
                    "INSERT INTO members (member_id, member_no, full_name, display_name, abbreviation, group_name,"
                    " organization, file_storage_location, is_proprietary) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (mid, f"{index:03d}", full, disp, abbr, group, org, storage, prop),
                )
                db.conn.execute(
                    "INSERT INTO project_members (project_id, member_id, assignment_name, is_proprietary,"
                    " include_in_monthly, start_date) VALUES (?, ?, ?, ?, ?, ?)",
                    (PROJECT_ID, mid, f"{index:02d}_{disp}", prop, monthly, start.isoformat()),
                )
        for day, name in PUBLIC_HOLIDAYS:
            db.add_holiday(name, '祝日', day)
        for name, category, first, last in COMPANY_HOLIDAYS:
            db.add_holiday(name, category, first, last)
        year, month = months[1]
        db.add_holiday('PJ定例（全体会）', 'イベント', f"{year}-{month:02d}-15", project_id=PROJECT_ID)


def _create_schedule(path: Path, display: str, group: str, months, base: date,
                     rng: random.Random, use_formula_dates: bool) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    holidays = {date.fromisoformat(d) for d, _ in PUBLIC_HOLIDAYS}
    for year, month in months:
        ws = wb.create_sheet(to_yyyymm(year, month))
        ws['A1'] = f"{display} 行動予定・実績 {year}年{month}月"
        ws['G1'] = '実績'
        ws['G1'].fill = ACTUAL_FILL
        for col, title in enumerate(SCHEDULE_HEADERS, start=1):
            cell = ws.cell(2, col, title)
            cell.font = Font(bold=True)
            cell.fill = ACTUAL_FILL if 7 <= col <= 12 else HEADER_FILL
            cell.border = THIN_BORDER
        start, end = month_range(year, month)
        for row, day in enumerate(iter_dates(start, end), start=3):
            if use_formula_dates and row > 3:
                ws.cell(row, 1, f"=A{row - 1}+1")
            else:
                ws.cell(row, 1, datetime(day.year, day.month, day.day))
            ws.cell(row, 1).number_format = 'm/d(aaa)'
            _fill_day(ws, row, day, group, day in holidays, day <= base, rng)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _fill_day(ws, row: int, day: date, group: str, is_holiday: bool, past: bool, rng: random.Random) -> None:
    off = day.weekday() >= 5 or is_holiday
    roll = rng.random()
    if off:
        if roll < 0.05:
            attendance, location = '休日出勤', group
        else:
            return
    elif roll < 0.06:
        attendance, location = '休暇', ''
    elif roll < 0.09:
        attendance, location = rng.choice(('A休', 'P休')), group
    elif roll < 0.25:
        attendance, location = '在宅', '自宅'
    elif roll < 0.33:
        attendance, location = '出張', '客先（架空商事）'
    else:
        attendance, location = '出社', group

    ws.cell(row, 2, attendance)
    ws.cell(row, 3, location)
    ws.cell(row, 4, '〇' if attendance in ('在宅', '出張') else '')
    ws.cell(row, 5, '〇' if attendance == '出張' and rng.random() < 0.5 else '')
    ws.cell(row, 6, '〇' if attendance == '在宅' else '')
    if past and attendance != '休暇':
        total = 4.0 if attendance in ('A休', 'P休') else rng.choice((7.5, 8.0, 8.0, 8.5, 9.0, 10.0))
        split = rng.sample(range(7, 12), 2)
        ws.cell(row, split[0], total - 2.0 if total > 2 else total)
        if total > 2:
            ws.cell(row, split[1], 2.0)
        if rng.random() < 0.1:
            ws.cell(row, 12, 1.0)
    if rng.random() < 0.05:
        ws.cell(row, 13, '架空のメモ')


def _create_templates(directory: Path) -> None:
    weekdays = ('月', '火', '水', '木', '金', '土', '日')

    wb = Workbook()
    ws = wb.active
    ws.title = 'yyyymm'
    for col, name in enumerate(weekdays, start=2):
        cell = ws.cell(1, col, name)
        cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, THIN_BORDER
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 22
    for week in range(6):
        ws.row_dimensions[3 + week * 2].height = 90
        for col in range(2, 9):
            ws.cell(2 + week * 2, col).border = THIN_BORDER
            ws.cell(3 + week * 2, col).border = THIN_BORDER
    wb.save(directory / '月間カレンダーテンプレ.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = 'yyyymm'
    ws['A1'] = '一覧カレンダー'
    ws['A1'].font = Font(bold=True, size=14)
    for col, (name, width) in enumerate(
            (('日付', 12), ('拠点', 12), ('メンバ', 12), ('勤怠', 10), ('行先', 20),
             ('PC持出', 8), ('宿泊', 8), ('wifi', 8), ('備考', 24)), start=1):
        cell = ws.cell(2, col, name)
        cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = width
    wb.save(directory / '一覧カレンダーテンプレ.xlsx')

    for suffix, title in (('自社向け', '作業実績表（自社向け）'), ('PJ向け', '作業実績表（PJ向け）')):
        wb = Workbook()
        ws = wb.active
        ws.title = 'xx月実績'
        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=14)
        for ref, label in (('A7', '年'), ('D7', '月'), ('G7', '所属'), ('M7', '氏名'),
                           ('E9', 'PJ名'), ('E10', 'PJコード')):
            ws[ref] = label
            ws[ref].font = Font(bold=True)
        for col, label in ((2, '日'), (4, '曜日'), (6, '外部設計'), (7, '内部設計'),
                           (8, '製造/単体テスト'), (9, '会議'), (10, 'その他')):
            cell = ws.cell(13, col, label)
            cell.font, cell.fill, cell.border = Font(bold=True), HEADER_FILL, THIN_BORDER
        for row in range(14, 45):
            for col in (2, 4, 6, 7, 8, 9, 10):
                ws.cell(row, col).border = THIN_BORDER
        # 本番テンプレートと同じ要素（画像・合計の数式・印刷範囲）を持たせ、書き込みで壊れないことを試せるようにする
        ws['B45'] = '合計'
        for col in 'FGHIJ':
            ws[f'{col}45'] = f'=SUM({col}14:{col}44)'
        ws.print_area = 'A1:N45'
        ws.add_image(_sample_logo(), 'K1')
        wb.save(directory / f'作業実績表テンプレート_{suffix}.xlsx')


def _sample_logo():
    """架空のロゴ画像（Pillow で描く）"""
    from io import BytesIO

    from openpyxl.drawing.image import Image as XlImage
    from PIL import Image, ImageDraw

    image = Image.new('RGB', (160, 48), '#2F5597')
    ImageDraw.Draw(image).rectangle((6, 6, 42, 42), fill='#FFFFFF')
    buffer = BytesIO()
    image.save(buffer, format='PNG')
    buffer.seek(0)
    return XlImage(buffer)
