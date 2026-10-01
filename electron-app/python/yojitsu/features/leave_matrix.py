"""休暇ステータス一覧（期間 × メンバ）

記号: 祝日=□ / 休暇=〇 / A休=△ / P休=▽ / 出社・その他=空欄
"""

import logging
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..context import RunContext
from ..dates import iter_dates, iter_year_months
from ..excel import THIN_BORDER, save_with_fallback, solid
from ..members import sort_by_group

logger = logging.getLogger(__name__)

LEGEND = '凡例: 空欄=出社  □=祝日  〇=休暇  △=A休  ▽=P休'
SATURDAY_FILL = solid('DDEBF7')
HOLIDAY_FILL = solid('FBE5D6')
LEGEND_FILL = solid('FFF2CC')
CENTER = Alignment(horizontal='center', vertical='center')


def symbol_for(attendance: str, is_public_holiday: bool) -> str:
    # 「A休暇」のような表記もあるため半休を先に判定する
    if 'A休' in attendance:
        return '△'
    if 'P休' in attendance:
        return '▽'
    if '休暇' in attendance:
        return '〇'
    return '□' if is_public_holiday else ''


def run(ctx: RunContext, start: date, end: date) -> dict:
    if start > end:
        raise ValueError('休暇一覧の開始日が終了日より後です')

    holidays = ctx.holidays(start, end)
    members = sort_by_group(ctx.members(start, end), ctx.site)
    days = list(iter_dates(start, end))

    attendance = {}
    for member in members:
        if not ctx.schedules.has_file(member.file_name):
            ctx.warn('個別予定ファイルが見つかりません', member.display_name)
        by_day = {}
        for y, m in iter_year_months(start, end):
            for day, entry in (ctx.schedules.month(member.file_name, y, m) or {}).items():
                # 土日祝の「休暇」は予定なしと同じ
                skip = entry.is_vacation and not holidays.is_business_day(day)
                by_day[day] = '' if skip else entry.attendance
        attendance[member.member_id] = by_day

    wb = Workbook()
    ws = wb.active
    ws.title = f"{start:%Y%m%d}_{end:%Y%m%d}"
    last_col = 1 + len(days)

    legend = ws.cell(1, 1, LEGEND)
    legend.font = Font(bold=True)
    legend.fill = LEGEND_FILL
    if last_col >= 2:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    header = ws.cell(2, 1, 'メンバ')
    header.alignment, header.font = CENTER, Font(bold=True)
    for col, day in enumerate(days, start=2):
        cell = ws.cell(2, col, day)
        cell.number_format = 'm/d'
        cell.alignment, cell.font = CENTER, Font(bold=True)
        if day.weekday() == 6 or holidays.is_public_holiday(day):
            cell.fill = HOLIDAY_FILL
        elif day.weekday() == 5:
            cell.fill = SATURDAY_FILL

    for row, member in enumerate(members, start=3):
        ws.cell(row, 1, member.display_name)
        for col, day in enumerate(days, start=2):
            is_holiday = holidays.is_public_holiday(day)
            mark = symbol_for(attendance[member.member_id].get(day, ''), is_holiday)
            cell = ws.cell(row, col, mark)
            cell.alignment = CENTER
            if mark or is_holiday or day.weekday() == 6:
                cell.fill = HOLIDAY_FILL
            elif day.weekday() == 5:
                cell.fill = SATURDAY_FILL

    for row in ws.iter_rows(min_row=1, max_row=2 + len(members), max_col=last_col):
        for cell in row:
            cell.border = THIN_BORDER
    ws.freeze_panes = 'B3'
    ws.column_dimensions['A'].width = 18
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 6.5

    output = ctx.output_dir() / f"休暇ステータス一覧_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    saved = save_with_fallback(wb, output)
    logger.info(f"休暇ステータス一覧を出力: {saved}")
    return {'outputs': [str(saved)]}
