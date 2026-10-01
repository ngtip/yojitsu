"""架空データで同期→全帳票作成までを通す"""

from pathlib import Path

from openpyxl import load_workbook

from yojitsu.schedule import read_month


def _sync(cli, db_file):
    code, result = cli('sync', '--db-file', db_file)
    assert code == 0, result
    return result


def test_sync_excludes_external_members(cli, db_file, sample_env):
    result = _sync(cli, db_file)
    names = sorted(Path(p).name for p in result['outputs'])
    assert names == ['01_甲野.xlsx', '02_乙川.xlsx', '03_丙田.xlsx', '04_丁村.xlsx', '05_戊井.xlsx']
    assert result['data']['downloaded_count'] == 5


def test_sync_reports_missing_remote(cli, db_file):
    cli('db', '--db-file', db_file, '--action', 'set-config',
        '--tool-name', 'sharepoint', '--config-key', 'remote_work_dir', '--config-value', '/sites/none')
    code, result = cli('sync', '--db-file', db_file)
    assert code == 1 and not result['success']
    assert 'リモートフォルダが見つかりません' in result['error']


def test_all_tasks(cli, db_file, sample_env):
    _sync(cli, db_file)
    for task in ('monthly-calendar', 'list-calendar', 'hours-summary', 'submit-files'):
        code, result = cli('run', task, '--db-file', db_file, '--year-month', '202610')
        assert code == 0, result
        assert result['outputs'] and all(Path(p).exists() for p in result['outputs'])

    code, result = cli('run', 'leave-matrix', '--db-file', db_file,
                       '--start-date', '2026-09-01', '--end-date', '2026-10-31')
    assert code == 0, result
    assert Path(result['outputs'][0]).name == '休暇ステータス一覧.xlsx'
    ws = load_workbook(result['outputs'][0]).active
    assert ws.title == '20260901_20261031'
    assert ws.max_column == 1 + 61
    assert ws.cell(3, 1).value == '甲野'


def test_calendars_hold_two_months_newest_first(cli, db_file):
    _sync(cli, db_file)
    for task, name in (('monthly-calendar', '月間カレンダー.xlsx'), ('list-calendar', '一覧カレンダー.xlsx')):
        _, result = cli('run', task, '--db-file', db_file, '--year-month', '202610')
        wb = load_workbook(result['outputs'][0])
        assert wb.sheetnames == ['202611', '202610']
        assert wb.active.title == '202610'


def test_list_calendar_weekend_shows_only_scheduled_members(cli, db_file, sample_env):
    _sync(cli, db_file)
    _, result = cli('run', 'list-calendar', '--db-file', db_file, '--year-month', '202610')
    ws = load_workbook(result['outputs'][0])['202610']
    rows_by_day = {}
    for row in range(3, ws.max_row + 1):
        day = ws.cell(row, 1).value.date()
        rows_by_day.setdefault(day, []).append([ws.cell(row, c).value for c in range(3, 5)])
    weekday_rows = rows_by_day[next(d for d in rows_by_day if d.weekday() < 5)]
    assert len(weekday_rows) == 5                       # 予定ファイルがある5人
    for day, rows in rows_by_day.items():
        if day.weekday() >= 5:
            assert all(r[1] for r in rows) or rows == [[None, None]]


def test_submit_files_match_schedule(cli, db_file, sample_env):
    _sync(cli, db_file)
    _, result = cli('run', 'submit-files', '--db-file', db_file, '--year-month', '202610')
    outputs = [Path(p) for p in result['outputs']]
    # BP3人は2種、プロパー2人はPJ向けのみ
    assert len([p for p in outputs if p.parent.name == '自社向け']) == 3
    assert len([p for p in outputs if p.parent.name == 'PJ向け']) == 5

    own = next(p for p in outputs if p.parent.name == '自社向け' and '丙田' in p.name)
    pj = next(p for p in outputs if p.parent.name == 'PJ向け' and '丙田' in p.name)
    assert own.name == 'サンプルPJ（架空SYS_丙田）作業実績表_202610.xlsx'
    assert pj.name == 'サンプルPJ（自社_丙田）作業実績表_202610.xlsx'

    schedule = read_month(sample_env / 'schedules' / '03_丙田.xlsx', 2026, 10)
    expected = sum(e.project_hours for e in schedule.values())
    own_ws = load_workbook(own)['10月実績']
    pj_ws = load_workbook(pj)['10月実績']
    own_total = sum(own_ws.cell(r, c).value or 0 for r in range(14, 45) for c in range(6, 11))
    pj_total = sum(pj_ws.cell(r, 6).value or 0 for r in range(14, 45))
    assert own_total == pj_total == expected > 0
    assert own_ws['H7'].value == '架空システム株式会社'
    assert pj_ws['H7'].value == 'サンプル自社株式会社'
    assert own_ws['F10'].value == 'S-0001' and own_ws['F10'].number_format == '@'
    assert own_ws['B44'].value == 31 and own_ws['D14'].value == '木'


def test_db_actions_for_settings_screen(cli, db_file):
    code, result = cli('db', '--db-file', db_file, '--action', 'get-project-members',
                       '--project-id', 'PJ-SAMPLE', '--as-of-date', '2026-10-01')
    assert code == 0 and result['data']['count'] == 5      # 戊井は 10/16 から参画

    code, result = cli('db', '--db-file', db_file, '--action', 'add-holiday',
                       '--name', '架空の日', '--category', '祝日', '--start-date', '2026-10-30')
    assert code == 0 and result['data']['holiday_id']

    code, result = cli('db', '--db-file', db_file, '--action', 'update-member', '--member-id', 'M001',
                       '--is-proprietary', '2')
    assert code == 1 and 'is_proprietary' in result['error']


def test_missing_template_is_reported(cli, db_file, sample_env):
    (sample_env / 'templates' / '月間カレンダーテンプレ.xlsx').unlink()
    code, result = cli('run', 'monthly-calendar', '--db-file', db_file, '--year-month', '202610')
    assert code == 1 and 'テンプレートが見つかりません' in result['error']


def test_hours_summary_marks_partial_participation(cli, db_file):
    """途中離脱→再参画（丁村）と途中参画（戊井）: 期間外は予測に入れず、期間外の実績を別列に出す"""
    _sync(cli, db_file)
    code, result = cli('run', 'hours-summary', '--db-file', db_file, '--year-month', '202610')
    assert code == 0
    ws = load_workbook(result['outputs'][0]).active
    rows = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(1, 14)] for r in range(5, ws.max_row + 1)}
    assert len(rows) == 6                                   # 丁村は2回参画でも1行
    full, tei, bo = rows['甲野'], rows['丁村'], rows['戊井']
    # 2026年10月の営業日: 平日22日 - スポーツの日 = 21日
    assert full[11] == '全期間' and full[7] == 21
    assert tei[11] == '10/1〜10/10、10/21〜10/31' and tei[7] == 7 + 8
    assert bo[11] == '10/16〜10/31' and bo[7] == 11
    assert tei[12] > 0 and bo[12] > 0 and full[12] == 0     # 期間外の実績（合計には含む）
    warned = {w['member'] for w in result['warnings'] if '参画期間外' in w['message']}
    assert warned == {'丁村', '戊井'}


def test_leave_matrix_is_overwritten_with_latest_period(cli, db_file):
    _sync(cli, db_file)
    for start, end in (('2026-09-01', '2026-10-31'), ('2026-10-01', '2026-10-15')):
        code, result = cli('run', 'leave-matrix', '--db-file', db_file, '--start-date', start, '--end-date', end)
        assert code == 0
    wb = load_workbook(result['outputs'][0])
    assert wb.sheetnames == ['20261001_20261015'] and wb.active.max_column == 1 + 15
